import base64
import copy
import io
import re
import unittest

from docx import Document
from PIL import Image
from pypdf import PdfReader

from services.cv_export import TEMPLATE_IDS, render, photo_stream, plain_text, visual_sections
from services.cv_builder import filename


def sample(language='en', count=1):
    return dict(full_name='Анна Example' if language == 'ru' else 'Anna Example',
                target_role='Data Analyst', email='anna@example.com', current_location='Dubai',
                summary='Prepared operational reports and documented data checks.', skills='SQL, Python, Power BI',
                languages='English C1; Русский — родной',
                experience=[dict(company=f'Example Company {i}', role='Data Analyst', location='Dubai',
                                 start_date='2020-01', end_date='Present',
                                 description='Prepared reports and documented SQL queries. ' * 8 + f' ItemEnd{i}') for i in range(count)],
                education=[dict(institution='Example University', qualification='BSc', field='Information Systems', dates='2016–2020')])


class CVTemplateTests(unittest.TestCase):
    def test_presentation_does_not_mutate_data_or_active_cv_text(self):
        data = sample('ru', 3)
        original = copy.deepcopy(data)
        evidence = plain_text(data, 'ru')
        for template in TEMPLATE_IDS:
            for extension in ('pdf', 'docx'):
                render(data, extension, 'ru', template)
                self.assertEqual(data, original)
                self.assertEqual(plain_text(data, 'ru'), evidence)

    def test_role_company_dates_are_distinct_without_invented_fields(self):
        data = {'experience': [{'company': 'Only known company'},
                               {'role': 'Engineer', 'start_date': '2020', 'end_date': 'Present'}]}
        parts = visual_sections(data, 'ru')['experience']
        self.assertIn(('organization', 'Only known company'), parts)
        self.assertIn(('entry_title', 'Engineer'), parts)
        self.assertIn(('meta', '2020 - По настоящее время'), parts)
        self.assertNotIn('Dubai', str(parts))

    def test_medium_demo_one_page_and_complete(self):
        from examples.cv_redesign.generate import sample as demo
        for template in TEMPLATE_IDS:
            for language in ('ru', 'en'):
                data = demo(language)
                pdf = PdfReader(io.BytesIO(render(data, 'pdf', language, template)))
                self.assertEqual(len(pdf.pages), 1)
                self.assertIn('Microsoft Power BI Data Analyst', pdf.pages[0].extract_text())

    def test_long_identity_wraps_and_is_not_lost(self):
        data = sample()
        data['full_name'] = 'Alexandria ' * 25
        data['target_role'] = 'Senior operations and reporting specialist ' * 10
        for template in TEMPLATE_IDS:
            pdf = PdfReader(io.BytesIO(render(data, 'pdf', 'en', template)))
            extracted = re.sub(r'\s', '', ''.join(page.extract_text() for page in pdf.pages))
            self.assertEqual(extracted.count('Alexandria'), 25)
            self.assertEqual(extracted.count('Senioroperationsandreportingspecialist'), 10)
            self.assertIn('ItemEnd0', extracted)

    def test_no_photo_has_no_placeholder_and_modern_full_width_header(self):
        for template in TEMPLATE_IDS:
            doc = Document(io.BytesIO(render(sample(), 'docx', 'en', template)))
            self.assertEqual(len(doc.element.xpath('//w:drawing')), 0)
            if template == 'template_1':
                self.assertEqual(len(doc.tables[0]._tbl.tr_lst[0].tc_lst), 1)

    def test_long_docx_continues_outside_sidebar_table(self):
        for template in ('template_1', 'template_2'):
            doc = Document(io.BytesIO(render(sample('ru', 7), 'docx', 'ru', template)))
            self.assertIn('ItemEnd6', '\n'.join(p.text for p in doc.paragraphs))
            for p in doc.element.xpath('//w:p[w:pPr/w:pStyle[@w:val="Heading1"]]'):
                self.assertTrue(p.xpath('./w:pPr/w:keepNext'))

    def test_all_templates_ru_en_pdf_and_editable_docx(self):
        for template in TEMPLATE_IDS:
            for language in ('en', 'ru'):
                with self.subTest(template=template, language=language):
                    data = sample(language)
                    pdf = PdfReader(io.BytesIO(render(data, 'pdf', language, template)))
                    extracted = '\n'.join(p.extract_text() for p in pdf.pages)
                    doc = Document(io.BytesIO(render(data, 'docx', language, template)))
                    doc_text = '\n'.join(doc.element.xpath('//w:t/text()'))
                    for output in (extracted, doc_text):
                        self.assertIn(data['full_name'], re.sub(r'\s+', ' ', output))
                        self.assertIn('Example Company', output)
                        self.assertIn('Русский', output)
                        self.assertIn('Навыки' if language == 'ru' else 'Skills', output)
                    self.assertAlmostEqual(float(pdf.pages[0].mediabox.width), 595.28, delta=1)
                    self.assertAlmostEqual(doc.sections[0].page_width.mm, 210, delta=0.1)
                    self.assertEqual(len(doc.inline_shapes), 0)
                    if template == 'template_3':
                        self.assertEqual(len(doc.tables), 0)
                        self.assertLess(extracted.index('Work experience') if language == 'en' else extracted.index('Опыт работы'), extracted.index('Skills') if language == 'en' else extracted.index('Навыки'))

    def test_long_content_multiple_pages_and_complete_extraction(self):
        for template in TEMPLATE_IDS:
            data = sample('ru', 7)
            data['email'] = 'long.' + 'x'*180 + '@example.com'
            data['linkedin'] = 'https://www.linkedin.com/in/' + 'a'*350
            data['experience'][0]['company'] = 'VeryLongCompany' * 20
            data['experience'][0]['role'] = 'ExtremelyLongRole' * 15
            data['education'] *= 3
            pdf = PdfReader(io.BytesIO(render(data, 'pdf', 'ru', template)))
            self.assertGreaterEqual(len(pdf.pages), 2)
            output = ''.join(p.extract_text() for p in pdf.pages)
            self.assertIn('ItemEnd6', output)
            compact = re.sub(r'\s', '', output)
            self.assertIn(data['email'], compact)
            self.assertIn(data['linkedin'], compact)
            self.assertIn(data['experience'][0]['company'], compact)
            doc = Document(io.BytesIO(render(data, 'docx', 'ru', template)))
            doc_text = ''.join(doc.element.xpath('//w:t/text()'))
            self.assertIn('ItemEnd6', doc_text)
            self.assertIn(data['linkedin'], doc_text)

    def test_empty_sections_missing_photo_and_invalid_photo(self):
        for template in TEMPLATE_IDS:
            data = {'full_name': 'Test Name', 'target_role': 'Cook', 'experience': [{}], 'education': [{}], 'photo': 'invalid'}
            pdf = PdfReader(io.BytesIO(render(data, 'pdf', 'en', template)))
            text = ''.join(p.extract_text() for p in pdf.pages)
            for heading in ('Education', 'Work experience', 'Contact', 'References', 'Certifications'):
                self.assertNotIn(heading, text)
            doc = Document(io.BytesIO(render(data, 'docx', 'en', template)))
            self.assertEqual(len(doc.inline_shapes), 0)

    def test_portrait_crop_and_optional_photo_all_templates(self):
        image = io.BytesIO()
        Image.new('RGB', (50, 140), 'blue').save(image, 'PNG')
        data = sample()
        data['photo'] = base64.b64encode(image.getvalue()).decode()
        with Image.open(photo_stream(data)) as cropped:
            self.assertEqual(cropped.size, (400, 400))
        for template in TEMPLATE_IDS:
            doc = Document(io.BytesIO(render(data, 'docx', 'en', template)))
            self.assertEqual(len(doc.element.xpath('//w:drawing')), 1)
            if doc.inline_shapes:
                self.assertEqual(doc.inline_shapes[0].width, doc.inline_shapes[0].height)

    def test_filename_sanitization_retained(self):
        self.assertEqual(filename('Anna Example', 'pdf'), 'Anna_Example_CV.pdf')
        self.assertNotIn('/', filename('../Анна:Example', 'docx'))
