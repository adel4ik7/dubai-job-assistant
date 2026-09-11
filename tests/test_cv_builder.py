import io
import tempfile
import unittest
from pathlib import Path

from docx import Document
from PIL import Image
from pypdf import PdfReader

from db import Database
from services.cv_builder import CVBuilder, filename
from services.cv_export import plain_text, render
from services.product import UserFiles
from services.matcher import analyse_match


class CVBuilderTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.db = Database(self.root / 'test.db')
        self.db.upsert_user(1, 'test', 'Test')
        self.builder = CVBuilder(self.db, self.root / 'uploads')
        self.did = self.builder.create(1)

    def complete_basics(self):
        self.builder.set_section(1, self.did, 'basics', {'full_name': 'Анна Example', 'target_role': 'Data Analyst'})

    def test_draft_persistence_and_owner_isolation(self):
        self.assertEqual(self.builder.get(1, self.did)['data'], {})
        self.complete_basics()
        builder = CVBuilder(Database(self.db.path), self.builder.root)
        self.assertEqual(builder.get(1, self.did)['data']['full_name'], 'Анна Example')
        with self.assertRaises(ValueError):
            builder.get(2, self.did)
        with self.assertRaises(ValueError):
            builder.delete(2, self.did)
        self.assertEqual(builder.list(2), [])

    def test_multiple_entries_edit_and_remove(self):
        for i in range(2):
            self.builder.set_section(1, self.did, 'experience', {'company': f'Company {i}', 'role': 'Analyst', 'end_date': 'Present'})
            self.builder.set_section(1, self.did, 'education', {'institution': f'School {i}', 'qualification': 'BSc'})
        self.builder.set_section(1, self.did, 'experience', {'company': 'Edited'}, 0)
        data = self.builder.get(1, self.did)['data']
        self.assertEqual(len(data['experience']), 2)
        self.assertEqual(len(data['education']), 2)
        self.assertEqual(data['experience'][0]['company'], 'Edited')
        self.builder.remove_entry(1, self.did, 'education', 0)
        self.assertEqual(len(self.builder.get(1, self.did)['data']['education']), 1)

    def test_exports_ru_en_and_no_fabricated_fields(self):
        self.complete_basics()
        self.builder.set_section(1, self.did, 'skills', {'skills': 'SQL, Python'})
        for language, heading in [('ru', 'Навыки'), ('en', 'Skills')]:
            docx = self.builder.attachment(1, self.did, 'docx', language)
            text = '\n'.join(p.text for p in Document(io.BytesIO(docx.content)).paragraphs)
            pdf = self.builder.attachment(1, self.did, 'pdf', language)
            pdf_text = '\n'.join(page.extract_text() for page in PdfReader(io.BytesIO(pdf.content)).pages)
            for result in (text, pdf_text):
                self.assertIn('Анна Example', result)
                self.assertIn(heading, result)
                self.assertIn('SQL, Python', result)
                self.assertNotIn('Experience', result)
                self.assertNotIn('Present', result)
                self.assertNotIn('Nationality', result)
        self.assertEqual(docx.filename, 'Анна_Example_CV.docx')

    def test_active_snapshot_refresh_clone_and_delete(self):
        self.complete_basics()
        self.builder.set_section(1, self.did, 'skills', {'skills': 'SQL'})
        rid = self.builder.activate(1, self.did)
        active = self.db.active_resume(1)
        old_path = Path(active['file_path'])
        self.assertEqual(active['id'], rid)
        self.assertTrue(old_path.is_file())
        self.assertGreater(analyse_match(active['extracted_text'], 'SQL required')['score'], 0)
        self.builder.set_section(1, self.did, 'skills', {'skills': 'Python'})
        self.assertIn('SQL', self.db.active_resume(1)['extracted_text'])
        self.assertEqual(self.builder.activate(1, self.did), rid)
        self.assertFalse(old_path.exists())
        self.assertIn('Python', self.db.active_resume(1)['extracted_text'])
        clone = self.builder.duplicate(1, self.did)
        self.assertIsNone(self.builder.get(1, clone)['resume_id'])
        self.builder.delete(1, self.did)
        self.assertIsNone(self.db.active_resume(1))
        self.assertEqual(len(self.builder.list(1)), 1)

    def test_filename_sanitization(self):
        self.assertEqual(filename('FirstName LastName', 'pdf'), 'FirstName_LastName_CV.pdf')
        for name in ['../../CON:bad\\path', 'a\x00b', '   ', 'X' * 500, 'Адель Тест']:
            output = filename(name, 'docx')
            self.assertEqual(Path(output).name, output)
            self.assertNotIn('..', output)
            self.assertNotRegex(output, r'[<>:"/\\|?*\x00]')
            self.assertLess(len(output), 110)

    def test_photo_and_privacy_deletion(self):
        self.complete_basics()
        photo = io.BytesIO()
        Image.new('RGB', (50, 80), 'blue').save(photo, 'PNG')
        self.builder.set_section(1, self.did, 'photo', photo.getvalue())
        self.assertTrue(self.builder.get(1, self.did)['data']['photo'])
        self.assertEqual(len(Document(io.BytesIO(self.builder.attachment(1, self.did, 'docx').content)).inline_shapes), 1)
        self.assertTrue(self.builder.attachment(1, self.did, 'pdf').content.startswith(b'%PDF'))
        self.builder.activate(1, self.did)
        UserFiles(self.db, self.builder.root).delete_all(1)
        self.assertEqual(self.builder.list(1), [])
        self.assertEqual(list(self.builder.root.iterdir()), [])

    def test_validation_and_template_placeholder(self):
        with self.assertRaises(ValueError):
            self.builder.attachment(1, self.did, 'pdf')
        with self.assertRaises(ValueError):
            self.builder.set_section(1, self.did, 'basics', {'email': 'invalid'})
        with self.assertRaises(ValueError):
            self.builder.set_section(1, self.did, 'experience', {'start_date': '2025-12', 'end_date': '2024'})
        with self.assertRaises(ValueError):
            render({}, 'pdf', template='template_2')
        with self.assertRaises(ValueError):
            render({}, 'pdf', template='../secret')
        self.assertEqual(plain_text({}), '')
